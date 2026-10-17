"""Renders every slide of a .pptx through PowerPoint itself (COM automation),
one transparent layer per top-level shape, so the web player can animate
shapes individually while the final composite is pixel-identical to
PowerPoint's own rendering (fonts, spacing, freeforms, crops, effects).

Each layer's alpha is recovered by rendering it over white and over black
(difference matting): for a pixel composited as a*C + (1-a)*bg,
  white - black = (1 - a) * 255  ->  a,  and  black = a*C  ->  C.
Shapes whose (padded, rotated) bounds don't overlap share a render pass,
which cuts the number of exports several-fold.
"""
import math
import os
import shutil
import sys
import tempfile

import numpy as np
from PIL import Image

MSO_TRUE, MSO_FALSE = -1, 0
FILL_PICTURE = 6
TYPE_AUTOSHAPE, TYPE_FREEFORM, TYPE_GROUP = 1, 5, 6
TYPE_PICTURE, TYPE_PLACEHOLDER, TYPE_TEXTBOX = 13, 14, 17
FILLABLE_TYPES = {TYPE_AUTOSHAPE, TYPE_FREEFORM, TYPE_PLACEHOLDER, TYPE_TEXTBOX}
ALPHA_EPS = 2 / 255
# Margin around each shape's geometry. Anything rendering further out
# (shadows, glows, overflowing text) is caught by the edge check and re-rendered.
PAD = 0.008
TEXT_PAD_Y = 0.035


class PowerPointUnavailable(RuntimeError):
    pass


def powerpoint_available():
    """Cheap check (no PowerPoint launch): Windows, pywin32 importable, and
    PowerPoint registered for automation."""
    if sys.platform != "win32":
        return False
    try:
        import win32com.client  # noqa: F401
        import winreg
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"PowerPoint.Application\CLSID"))
        return True
    except Exception:
        return False


def _powerpoint_running():
    import subprocess
    out = subprocess.run(["tasklist", "/FI", "IMAGENAME eq POWERPNT.EXE", "/NH"],
                         capture_output=True, text=True, creationflags=0x08000000).stdout
    return "POWERPNT.EXE" in out.upper()


BUSY_HRESULTS = {-2147418111, -2147417846}  # RPC_E_CALL_REJECTED, RPC_E_SERVERCALL_RETRYLATER


def _com(fn, attempts=300, wait=0.2):
    """Run a COM call, retrying while PowerPoint reports itself busy
    (starting up, rendering, or serving another automation client)."""
    import time
    for i in range(attempts):
        try:
            return fn()
        except Exception as e:
            hresult = getattr(e, "hresult", None) or (e.args[0] if getattr(e, "args", None) else None)
            if hresult not in BUSY_HRESULTS or i == attempts - 1:
                raise
            time.sleep(wait)


def _open_app():
    try:
        import pythoncom
        import win32com.client
    except ImportError as e:
        raise PowerPointUnavailable("pywin32 is not installed") from e
    pythoncom.CoInitialize()
    try:
        # Early-bound wrappers: late-bound property access costs a
        # name lookup per call, which dominated render time.
        return win32com.client.gencache.EnsureDispatch("PowerPoint.Application")
    except Exception:
        try:
            return win32com.client.Dispatch("PowerPoint.Application")
        except Exception as e:
            raise PowerPointUnavailable(f"PowerPoint could not be started: {e}") from e


def _safe(fn, default=None):
    try:
        return fn()
    except Exception:
        return default


def _shape_info(shape):
    """Kind, text and geometry in one pass. Only properties valid for the
    shape's type are touched: COM exceptions are slow, so probing optional
    properties with try/except on every shape is expensive."""
    stype = shape.Type
    info = {
        "type": stype, "kind": "shape", "text": "",
        "left": shape.Left, "top": shape.Top, "right": shape.Left + shape.Width,
        "bottom": shape.Top + shape.Height, "rotation": shape.Rotation or 0,
    }
    if shape.HasChart == MSO_TRUE:
        info["kind"] = "chart"
        return info
    if shape.HasTable == MSO_TRUE:
        info["kind"] = "table"
        return info
    contained = stype
    if stype == TYPE_PLACEHOLDER:
        contained = _safe(lambda: shape.PlaceholderFormat.ContainedType, stype)
    if shape.HasTextFrame == MSO_TRUE and shape.TextFrame.HasText == MSO_TRUE:
        tr = shape.TextFrame.TextRange
        info["text"] = (tr.Text or "").strip()
        info["kind"] = "text"
        # Text overflowing its frame renders outside the shape geometry.
        bl, bt = tr.BoundLeft, tr.BoundTop
        info["left"], info["top"] = min(info["left"], bl), min(info["top"], bt)
        info["right"] = max(info["right"], bl + tr.BoundWidth)
        info["bottom"] = max(info["bottom"], bt + tr.BoundHeight)
    if contained == TYPE_PICTURE or (stype in FILLABLE_TYPES and _safe(lambda: shape.Fill.Type == FILL_PICTURE, False)):
        info["kind"] = "image"
    elif stype == TYPE_GROUP:
        info["kind"] = "group"
    return info


EMU_PER_POINT = 12700


def _xml_shape_info(shape):
    """Same fields as _shape_info, read from the .pptx XML via python-pptx —
    milliseconds instead of a slow cross-process COM round trip per property.
    Text bounds aren't knowable without layout; text shapes get extra
    vertical padding instead and the edge check catches the rest."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.oxml.ns import qn

    hidden = shape._element.find(".//" + qn("p:cNvPr"))
    info = {
        "hidden": hidden is not None and hidden.get("hidden") in ("1", "true"),
        "kind": "shape", "text": "",
        "left": (shape.left or 0) / EMU_PER_POINT, "top": (shape.top or 0) / EMU_PER_POINT,
        "right": ((shape.left or 0) + (shape.width or 0)) / EMU_PER_POINT,
        "bottom": ((shape.top or 0) + (shape.height or 0)) / EMU_PER_POINT,
        "rotation": getattr(shape, "rotation", 0) or 0,
    }
    if getattr(shape, "has_chart", False):
        info["kind"] = "chart"
    elif getattr(shape, "has_table", False):
        info["kind"] = "table"
    elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE or (
        shape.shape_type != MSO_SHAPE_TYPE.GROUP
        and (shape._element.find(".//" + qn("a:blipFill")) is not None
             or shape._element.find(".//" + qn("p:blipFill")) is not None)):
        info["kind"] = "image"
    elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        info["kind"] = "group"
    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
        info["text"] = shape.text_frame.text.strip()
        if info["kind"] == "shape":
            info["kind"] = "text"
    return info


def _pixel_box(info, sx, sy, W, H, pad_x, pad_y):
    """Axis-aligned pixel box of the (possibly rotated) shape, padded."""
    left, top = info["left"] * sx, info["top"] * sy
    w, h = info["right"] * sx - left, info["bottom"] * sy - top
    rot = math.radians(info["rotation"])
    cx, cy = left + w / 2, top + h / 2
    cos, sin = abs(math.cos(rot)), abs(math.sin(rot))
    hw, hh = (w * cos + h * sin) / 2, (w * sin + h * cos) / 2
    x0 = max(0, int(math.floor(cx - hw - pad_x)))
    y0 = max(0, int(math.floor(cy - hh - pad_y)))
    x1 = min(W, int(math.ceil(cx + hw + pad_x)))
    y1 = min(H, int(math.ceil(cy + hh + pad_y)))
    return x0, y0, x1, y1


def _overlaps(a, b):
    return a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]


def _plan_passes(items):
    """Greedy grouping of shapes into passes of mutually non-overlapping boxes."""
    passes = []
    for item in items:
        for p in passes:
            if all(not _overlaps(item["box"], other["box"]) for other in p):
                p.append(item)
                break
        else:
            passes.append([item])
    return passes


def _load(path):
    return np.asarray(Image.open(path).convert("RGB"), dtype=np.float32)


def _save_layer(rgb, alpha, path, lossy):
    rgba = np.dstack([np.clip(rgb, 0, 255), np.clip(alpha * 255, 0, 255)]).astype(np.uint8)
    img = Image.fromarray(rgba, "RGBA")
    if lossy:
        img.save(path, "WEBP", quality=90, method=4)
    else:
        img.save(path, "WEBP", lossless=True, method=4)


def render_presentation(pptx_path, out_dir, width=1920, progress=None):
    """Returns the deck JSON consumed by the web player (mode "raster")."""
    if sys.platform != "win32":
        raise PowerPointUnavailable("exact rendering needs PowerPoint on Windows")
    # PowerPoint does all its file I/O in a private temp folder: Store/MSIX
    # Office installs can't open or write paths under %LOCALAPPDATA% (where
    # the desktop app keeps its data), but the temp folder always works.
    work = tempfile.mkdtemp(prefix="holodeck-")
    try:
        return _render(pptx_path, out_dir, width, progress, work)
    finally:
        shutil.rmtree(work, ignore_errors=True)


def _render(pptx_path, out_dir, width, progress, work):
    os.makedirs(out_dir, exist_ok=True)
    src_copy = os.path.join(work, "input" + os.path.splitext(pptx_path)[1])
    shutil.copyfile(pptx_path, src_copy)
    tmp_w = os.path.join(work, "_w.png")
    tmp_b = os.path.join(work, "_b.png")
    tmp_full = os.path.join(work, "_full.png")

    try:
        from pptx import Presentation as XmlPresentation
        xml_slides = list(XmlPresentation(pptx_path).slides)
    except Exception:
        xml_slides = []

    say = lambda msg: progress and progress(0.0, msg)
    say("BOOTING NATIVE RENDER CORE...")
    already_running = _safe(_powerpoint_running, True)
    app = _open_app()
    say("MOUNTING PRESENTATION...")
    pres = _com(lambda: app.Presentations.Open(src_copy, MSO_TRUE, MSO_TRUE, MSO_FALSE))
    try:
        sw, sh = _com(lambda: (pres.PageSetup.SlideWidth, pres.PageSetup.SlideHeight))
        W = int(width)
        H = int(round(W * sh / sw))
        sx, sy = W / sw, H / sh
        n_slides = _com(lambda: pres.Slides.Count)
        say(f"DECK MOUNTED · {n_slides} SLIDES · {W}x{H} RENDER TARGET")
        slides_out = []

        def set_visible(shape, on):
            _com(lambda: setattr(shape, "Visible", MSO_TRUE if on else MSO_FALSE))

        for si in range(1, n_slides + 1):
            slide = _com(lambda: pres.Slides(si))

            def export(path):
                _com(lambda: slide.Export(path, "PNG", W, H))

            def set_background(rgb):
                _com(lambda: setattr(slide.Background.Fill.ForeColor, "RGB", rgb))
            if _safe(lambda: slide.SlideShowTransition.Hidden == MSO_TRUE, False):
                continue
            prefix = f"s{si:02d}"
            report = lambda frac, msg: progress and progress((si - 1 + frac) / n_slides, msg)
            report(0.0, f"RENDERING SLIDE {si}/{n_slides}")

            shapes = _com(lambda: [slide.Shapes(i) for i in range(1, slide.Shapes.Count + 1)])
            # Both APIs list a slide's top-level shapes in z-order; if the
            # counts ever disagree, trust COM (slower but authoritative).
            xml_shapes = list(xml_slides[si - 1].shapes) if si - 1 < len(xml_slides) else []
            if len(xml_shapes) == len(shapes):
                infos = [_xml_shape_info(x) for x in xml_shapes]
                visible = [not i["hidden"] for i in infos]
            else:
                infos = [None] * len(shapes)
                visible = _com(lambda: [s.Visible == MSO_TRUE for s in shapes])

            # Base: slide background + master/layout graphics, no slide shapes.
            for s, vis in zip(shapes, visible):
                if vis:
                    set_visible(s, False)
            export(tmp_full)
            base_img = Image.open(tmp_full).convert("RGB")
            base_name = f"{prefix}_base.webp"
            base_img.save(os.path.join(out_dir, base_name), "WEBP", quality=92)

            # Isolate shapes on a flat background so they can be matted.
            _com(lambda: setattr(slide, "FollowMasterBackground", MSO_FALSE))
            _com(lambda: setattr(slide, "DisplayMasterShapes", MSO_FALSE))
            _com(lambda: slide.Background.Fill.Solid())

            items = []
            for z, (s, vis, info) in enumerate(zip(shapes, visible, infos)):
                if not vis:
                    continue
                if info is None:
                    info = _com(lambda: _shape_info(s))
                pad_x = PAD * W
                # Text frequently overflows its frame vertically by a line.
                pad_y = max(pad_x, TEXT_PAD_Y * H) if info["text"] else pad_x
                items.append({
                    "z": z, "shape": s, "box": _pixel_box(info, sx, sy, W, H, pad_x, pad_y),
                    "kind": info["kind"], "text": info["text"],
                })
            passes = _plan_passes(items)

            def render_pass(group):
                for it in group:
                    set_visible(it["shape"], True)
                set_background(0xFFFFFF)
                export(tmp_w)
                set_background(0x000000)
                export(tmp_b)
                for it in group:
                    set_visible(it["shape"], False)
                return _load(tmp_w), _load(tmp_b)

            layers = []

            def matte(it, wimg, bimg):
                x0, y0, x1, y1 = it["box"]
                wc, bc = wimg[y0:y1, x0:x1], bimg[y0:y1, x0:x1]
                alpha = np.clip(1.0 - (wc - bc).mean(axis=2) / 255.0, 0.0, 1.0)
                return alpha, bc

            def clipped(it, alpha):
                # Ink touching an inner edge of the crop means the shape renders
                # beyond its estimated bounds (shadow, glow, overflowing text).
                x0, y0, x1, y1 = it["box"]
                edges = []
                if y0 > 0: edges.append(alpha[0, :])
                if y1 < H: edges.append(alpha[-1, :])
                if x0 > 0: edges.append(alpha[:, 0])
                if x1 < W: edges.append(alpha[:, -1])
                return any(e.max() > 0.05 for e in edges)

            def save(it, alpha, bc):
                ys, xs = np.nonzero(alpha > ALPHA_EPS)
                if len(xs) == 0:
                    return
                ty0, ty1, tx0, tx1 = ys.min(), ys.max() + 1, xs.min(), xs.max() + 1
                a = alpha[ty0:ty1, tx0:tx1]
                rgb = bc[ty0:ty1, tx0:tx1] / np.maximum(a, 1e-3)[..., None]
                rgb[a <= ALPHA_EPS] = 0
                a = np.where(a <= ALPHA_EPS, 0.0, a)
                name = f"{prefix}_z{it['z']:03d}.webp"
                _save_layer(rgb, a, os.path.join(out_dir, name), lossy=it["kind"] == "image")
                x0, y0 = it["box"][:2]
                layers.append({
                    "z": it["z"], "src": name, "kind": it["kind"], "text": it["text"],
                    "x": round((x0 + tx0) / W * 100, 4), "y": round((y0 + ty0) / H * 100, 4),
                    "w": round((tx1 - tx0) / W * 100, 4), "h": round((ty1 - ty0) / H * 100, 4),
                })

            solo = 0
            for pi, group in enumerate(passes):
                report((pi + 1) / (len(passes) + 1), f"RENDERING SLIDE {si}/{n_slides} · PASS {pi + 1}/{len(passes)}")
                wimg, bimg = render_pass(group)
                mattes = [matte(it, wimg, bimg) for it in group]
                bleeding = [clipped(it, m[0]) for it, m in zip(group, mattes)]
                if len(group) == 1:
                    # Exports are full-slide, so a lone bleeding shape just
                    # needs a wider crop of the same render.
                    if bleeding[0]:
                        group[0]["box"] = (0, 0, W, H)
                        mattes[0] = matte(group[0], wimg, bimg)
                    save(group[0], *mattes[0])
                    continue
                if not any(bleeding):
                    for it, m in zip(group, mattes):
                        save(it, *m)
                    continue
                # Something rendered past its box, so it may also have leaked
                # into a neighbour's crop: redo every shape of this pass alone.
                for it, bleeds in zip(group, bleeding):
                    if bleeds:
                        it["box"] = (0, 0, W, H)
                    solo += 1
                    w1, b1 = render_pass([it])
                    save(it, *matte(it, w1, b1))

            layers.sort(key=lambda l: l["z"])
            thumb = base_img.convert("RGBA")
            for l in layers:
                if l["w"] * l["h"] > 0.85 * 100 * 100:
                    l["kind"] = "backdrop"
                layer_img = Image.open(os.path.join(out_dir, l["src"])).convert("RGBA")
                thumb.alpha_composite(layer_img, (round(l["x"] / 100 * W), round(l["y"] / 100 * H)))
            thumb = thumb.convert("RGB")
            thumb.thumbnail((480, 480))
            thumb_name = f"{prefix}_thumb.jpg"
            thumb.save(os.path.join(out_dir, thumb_name), "JPEG", quality=85)

            notes = _safe(lambda: slide.NotesPage.Shapes.Placeholders(2).TextFrame.TextRange.Text, "") or ""
            slides_out.append({
                "base": base_name, "thumb": thumb_name, "layers": layers,
                "notes": notes, "passes": len(passes), "soloRetries": solo,
            })

        return {
            "mode": "raster",
            "aspectRatio": round(W / H, 6),
            "renderWidth": W, "renderHeight": H,
            "slides": slides_out,
        }
    finally:
        _safe(lambda: pres.Close())
        # Only quit a PowerPoint we launched ourselves — quitting a shared
        # instance would kill the user's (or another job's) open session.
        if not already_running and _safe(lambda: app.Presentations.Count, 1) == 0:
            _safe(lambda: app.Quit())
