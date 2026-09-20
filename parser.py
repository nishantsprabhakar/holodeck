"""Parses a .pptx file into a JSON-serializable structure preserving
text, formatting, positions, images and backgrounds for web rendering."""
import os
import base64
from lxml import etree
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_FILL_TYPE, MSO_COLOR_TYPE
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn

EMU_PER_INCH = 914400

ALIGN_MAP = {
    PP_ALIGN.LEFT: "left",
    PP_ALIGN.CENTER: "center",
    PP_ALIGN.RIGHT: "right",
    PP_ALIGN.JUSTIFY: "justify",
}

ANCHOR_MAP = {
    MSO_ANCHOR.TOP: "top",
    MSO_ANCHOR.MIDDLE: "middle",
    MSO_ANCHOR.BOTTOM: "bottom",
}

PIE_CHART_TYPES = {
    XL_CHART_TYPE.PIE, XL_CHART_TYPE.PIE_EXPLODED,
    XL_CHART_TYPE.DOUGHNUT, XL_CHART_TYPE.DOUGHNUT_EXPLODED,
}
LINE_CHART_TYPES = {
    XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS,
    XL_CHART_TYPE.LINE_STACKED, XL_CHART_TYPE.LINE_STACKED_100,
    XL_CHART_TYPE.LINE_MARKERS_STACKED, XL_CHART_TYPE.LINE_MARKERS_STACKED_100,
    XL_CHART_TYPE.XY_SCATTER, XL_CHART_TYPE.XY_SCATTER_LINES,
    XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, XL_CHART_TYPE.XY_SCATTER_SMOOTH,
    XL_CHART_TYPE.XY_SCATTER_SMOOTH_NO_MARKERS,
}
HORIZONTAL_BAR_TYPES = {
    XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.BAR_STACKED, XL_CHART_TYPE.BAR_STACKED_100,
}

# Maps python-pptx's MSO_THEME_COLOR member name to the <a:clrScheme> child
# tag that holds its actual RGB value. BACKGROUND_*/TEXT_* follow the
# standard (unmapped) PowerPoint default; a theme with a custom <p:clrMap>
# override is a rare enough case to accept as a known gap.
THEME_SLOT_MAP = {
    "DARK_1": "dk1", "LIGHT_1": "lt1", "DARK_2": "dk2", "LIGHT_2": "lt2",
    "ACCENT_1": "accent1", "ACCENT_2": "accent2", "ACCENT_3": "accent3",
    "ACCENT_4": "accent4", "ACCENT_5": "accent5", "ACCENT_6": "accent6",
    "HYPERLINK": "hlink", "FOLLOWED_HYPERLINK": "folHlink",
    "BACKGROUND_1": "lt1", "TEXT_1": "dk1", "BACKGROUND_2": "lt2", "TEXT_2": "dk2",
}

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
THEME_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"


def _load_theme_map(prs):
    """Resolves the presentation's <a:clrScheme> once, so scheme-referenced
    colors (very common in professionally designed templates — "Accent 1"
    etc. rather than hardcoded RGB) can be looked up instead of silently
    dropped."""
    try:
        master = prs.slide_masters[0]
        theme_part = master.part.part_related_by(THEME_REL)
        root = etree.fromstring(theme_part.blob)
        ns = {"a": A_NS}
        clr_scheme = root.find(".//a:clrScheme", ns)
        result = {}
        if clr_scheme is not None:
            for child in clr_scheme:
                tag = etree.QName(child).localname
                srgb = child.find("a:srgbClr", ns)
                sys_clr = child.find("a:sysClr", ns)
                val = srgb.get("val") if srgb is not None else (sys_clr.get("lastClr") if sys_clr is not None else None)
                if val:
                    result[tag] = "#" + val.upper()
        return result
    except Exception:
        return {}


def _pct(value_emu, total_emu):
    if not total_emu:
        return 0.0
    return round((value_emu / total_emu) * 100, 4)


def _color_to_hex(color_format, theme=None):
    """Best-effort extraction of an RGB hex string from a font/fill color."""
    try:
        if color_format is None:
            return None
        if color_format.type is None:
            return None
        if color_format.type == MSO_COLOR_TYPE.SCHEME:
            theme_color = color_format.theme_color
            slot = THEME_SLOT_MAP.get(theme_color.name if theme_color else None)
            return (theme or {}).get(slot)
        rgb = color_format.rgb
        if rgb is None:
            return None
        return "#" + str(rgb)
    except Exception:
        return None


def _avg_hex(colors):
    vals = [c for c in colors if c]
    if not vals:
        return None
    rs = [int(c[1:3], 16) for c in vals]
    gs = [int(c[3:5], 16) for c in vals]
    bs = [int(c[5:7], 16) for c in vals]
    return "#%02X%02X%02X" % (sum(rs) // len(rs), sum(gs) // len(gs), sum(bs) // len(bs))


def _fill_color(shape, theme=None):
    """Best-effort fill color, approximating gradients/patterns with a
    representative solid so shapes aren't silently dropped."""
    try:
        fill = shape.fill
        ftype = fill.type
        if ftype is None:
            return None
        if ftype == MSO_FILL_TYPE.SOLID:
            return _color_to_hex(fill.fore_color, theme)
        if ftype == MSO_FILL_TYPE.GRADIENT:
            try:
                stops = list(fill.gradient_stops)
                # Average every stop rather than taking just the first —
                # gradients commonly start on a pale/washed-out tone (e.g.
                # #F3D9AE fading to #A9773E), so the endpoint alone looks
                # nothing like the shape's actual overall visual color.
                return _avg_hex([_color_to_hex(s.color, theme) for s in stops])
            except Exception:
                return None
        if ftype == MSO_FILL_TYPE.PATTERNED:
            return _color_to_hex(fill.fore_color, theme)
    except Exception:
        return None
    return None


def _line_info(shape, theme=None):
    """Stroke color/width, so outline-only shapes (unfilled boxes,
    connectors, lines) still render instead of vanishing entirely."""
    try:
        line = shape.line
        if line.color is None or line.color.type is None:
            return None
        color = _color_to_hex(line.color, theme)
        if not color:
            return None
        width_pt = 1.0
        try:
            if line.width is not None:
                width_pt = round(line.width.pt, 2)
        except Exception:
            pass
        return {"color": color, "widthPt": width_pt}
    except Exception:
        return None


def _slide_background(slide, prs, theme=None):
    try:
        bg = slide.background
        fill = bg.fill
        if fill.type == 1:
            hexval = _color_to_hex(fill.fore_color, theme)
            if hexval:
                return {"type": "color", "value": hexval}
    except Exception:
        pass
    return None


def _extract_runs(paragraph, theme=None):
    runs = []
    for run in paragraph.runs:
        font = run.font
        size_pt = None
        try:
            if font.size is not None:
                size_pt = round(font.size.pt, 1)
        except Exception:
            size_pt = None
        color = _color_to_hex(font.color, theme) if font.color is not None else None
        runs.append({
            "text": run.text,
            "bold": bool(font.bold) if font.bold is not None else False,
            "italic": bool(font.italic) if font.italic is not None else False,
            "underline": bool(font.underline) if font.underline is not None else False,
            "color": color,
            "font": font.name,
            "sizePt": size_pt,
        })
    if not runs and paragraph.text:
        runs.append({
            "text": paragraph.text,
            "bold": False, "italic": False, "underline": False,
            "color": None, "font": None, "sizePt": None,
        })
    return runs


def _extract_text_frame(shape, theme=None):
    paragraphs = []
    tf = shape.text_frame
    for para in tf.paragraphs:
        align = ALIGN_MAP.get(para.alignment, None)
        level = para.level or 0
        runs = _extract_runs(para, theme)
        if any(r["text"] for r in runs):
            paragraphs.append({"align": align, "level": level, "runs": runs})
    return paragraphs


def _autofit(shape):
    """Reads PowerPoint's own <a:normAutofit> shrink-to-fit values, if saved."""
    try:
        body_pr = shape.text_frame._txBody.bodyPr
        node = body_pr.find(qn("a:normAutofit"))
        if node is not None:
            fs = node.get("fontScale")
            lr = node.get("lnSpcReduction")
            font_scale = (int(fs) / 100000) if fs else 1.0
            ln_reduction = (int(lr) / 100000) if lr else 0.0
            return font_scale, ln_reduction
    except Exception:
        pass
    return 1.0, 0.0


def _text_frame_meta(shape, prs):
    tf = shape.text_frame
    try:
        anchor = ANCHOR_MAP.get(tf.vertical_anchor, "top")
    except Exception:
        anchor = "top"
    try:
        word_wrap = tf.word_wrap is not False
    except Exception:
        word_wrap = True
    try:
        margin_l = _pct(tf.margin_left if tf.margin_left is not None else 91440, prs.slide_width)
        margin_r = _pct(tf.margin_right if tf.margin_right is not None else 91440, prs.slide_width)
        margin_t = _pct(tf.margin_top if tf.margin_top is not None else 45720, prs.slide_height)
        margin_b = _pct(tf.margin_bottom if tf.margin_bottom is not None else 45720, prs.slide_height)
    except Exception:
        margin_l = margin_r = margin_t = margin_b = 0.0
    font_scale, ln_reduction = _autofit(shape)
    return {
        "anchor": anchor,
        "wordWrap": word_wrap,
        "marginLeftPct": margin_l,
        "marginRightPct": margin_r,
        "marginTopPct": margin_t,
        "marginBottomPct": margin_b,
        "fontScale": font_scale,
        "lnSpcReduction": ln_reduction,
    }


def _series_color(series, theme=None):
    try:
        fill = series.format.fill
        if fill.type == MSO_FILL_TYPE.SOLID:
            return _color_to_hex(fill.fore_color, theme)
    except Exception:
        pass
    return None


def _point_color(series, idx, theme=None):
    try:
        pt = series.points[idx]
        fill = pt.format.fill
        if fill.type == MSO_FILL_TYPE.SOLID:
            return _color_to_hex(fill.fore_color, theme)
    except Exception:
        pass
    return None


def _extract_chart(shape, theme=None):
    """Extracts chart type, categories and series data so charts render (and
    animate) as real web charts instead of being silently skipped."""
    try:
        chart = shape.chart
        ctype = chart.chart_type
    except Exception:
        return None

    if ctype in PIE_CHART_TYPES:
        kind = "pie"
    elif ctype in LINE_CHART_TYPES:
        kind = "line"
    else:
        kind = "bar"
    orientation = "horizontal" if ctype in HORIZONTAL_BAR_TYPES else "vertical"

    try:
        categories = [str(c) if c is not None else "" for c in chart.plots[0].categories]
    except Exception:
        categories = []

    series_out = []
    try:
        for series in chart.series:
            values = []
            for v in series.values:
                try:
                    values.append(float(v) if v is not None else 0.0)
                except Exception:
                    values.append(0.0)
            entry = {"name": series.name or "", "values": values}
            if kind == "pie":
                entry["colors"] = [_point_color(series, i, theme) for i in range(len(values))]
            else:
                entry["color"] = _series_color(series, theme)
            series_out.append(entry)
    except Exception:
        pass

    if not series_out:
        return None

    return {
        "chartType": kind,
        "orientation": orientation,
        "categories": categories,
        "series": series_out,
    }


def _save_image(shape, assets_dir, asset_prefix, counter):
    try:
        image = shape.image
        return _write_image_bytes(image.blob, image.ext, assets_dir, asset_prefix, counter)
    except Exception:
        return None


def _write_image_bytes(blob, ext, assets_dir, asset_prefix, counter):
    try:
        fname = f"{asset_prefix}_{counter}.{ext}"
        fpath = os.path.join(assets_dir, fname)
        with open(fpath, "wb") as f:
            f.write(blob)
        return fname
    except Exception:
        return None


def _picture_fill_image_part(shape):
    """A freeform/autoshape can use an image as its fill (a very common way
    to place cropped photos, icons and decorative graphics in PowerPoint).
    python-pptx's high-level FillFormat has no accessor for this, so read
    the <a:blipFill>/<a:blip> relationship directly from the shape XML."""
    try:
        sp_pr = shape._element.spPr
        blip_fill = sp_pr.find(qn("a:blipFill"))
        if blip_fill is None:
            return None
        blip = blip_fill.find(qn("a:blip"))
        if blip is None:
            return None
        r_id = blip.get(qn("r:embed"))
        if not r_id:
            return None
        return shape.part.related_part(r_id)
    except Exception:
        return None


def _save_picture_fill(shape, assets_dir, asset_prefix, counter):
    part = _picture_fill_image_part(shape)
    if part is None:
        return None
    try:
        image = part.image
        return _write_image_bytes(image.blob, image.ext, assets_dir, asset_prefix, counter)
    except Exception:
        return None


def _walk_shapes(shapes, prs, assets_dir, asset_prefix, counter_ref, theme=None):
    elements = []
    for shape in shapes:
        try:
            left = shape.left if shape.left is not None else 0
            top = shape.top if shape.top is not None else 0
            width = shape.width if shape.width is not None else 0
            height = shape.height if shape.height is not None else 0
            rect = {
                "leftPct": _pct(left, prs.slide_width),
                "topPct": _pct(top, prs.slide_height),
                "widthPct": _pct(width, prs.slide_width),
                "heightPct": _pct(height, prs.slide_height),
                "rotation": getattr(shape, "rotation", 0) or 0,
            }

            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                sub = _walk_shapes(shape.shapes, prs, assets_dir, asset_prefix, counter_ref, theme)
                elements.extend(sub)
                continue

            if getattr(shape, "has_chart", False):
                chart_data = _extract_chart(shape, theme)
                if chart_data:
                    elements.append({"type": "chart", **chart_data, **rect})
                continue

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                counter_ref[0] += 1
                fname = _save_image(shape, assets_dir, asset_prefix, counter_ref[0])
                if fname:
                    elements.append({"type": "image", "src": fname, **rect})
                continue

            try:
                shape_fill_type = shape.fill.type
            except Exception:
                shape_fill_type = None

            if shape_fill_type == MSO_FILL_TYPE.PICTURE:
                counter_ref[0] += 1
                fname = _save_picture_fill(shape, assets_dir, asset_prefix, counter_ref[0])
                if fname:
                    elements.append({"type": "image", "src": fname, **rect})
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        paragraphs = _extract_text_frame(shape, theme)
                        meta = _text_frame_meta(shape, prs)
                        elements.append({"type": "text", "fill": None, "stroke": None, "paragraphs": paragraphs, **meta, **rect})
                    continue

            if shape.has_table:
                table = shape.table
                rows = []
                for row in table.rows:
                    cells = []
                    for cell in row.cells:
                        cell_paras = []
                        for para in cell.text_frame.paragraphs:
                            runs = _extract_runs(para, theme)
                            if any(r["text"] for r in runs):
                                cell_paras.append({
                                    "align": ALIGN_MAP.get(para.alignment, None),
                                    "runs": runs,
                                })
                        cells.append(cell_paras)
                    rows.append(cells)
                elements.append({"type": "table", "rows": rows, **rect})
                continue

            fill_color = _fill_color(shape, theme)
            line_info = _line_info(shape, theme)
            has_text = shape.has_text_frame and shape.text_frame.text.strip()
            has_visible = fill_color or line_info

            if has_visible and not has_text:
                elements.append({"type": "shape", "fill": fill_color, "stroke": line_info, **rect})
            elif has_visible and has_text:
                paragraphs = _extract_text_frame(shape, theme)
                meta = _text_frame_meta(shape, prs)
                elements.append({"type": "text", "fill": fill_color, "stroke": line_info, "paragraphs": paragraphs, **meta, **rect})
            elif has_text:
                paragraphs = _extract_text_frame(shape, theme)
                meta = _text_frame_meta(shape, prs)
                is_title = False
                try:
                    is_title = shape.name.lower().startswith("title") or (
                        shape.placeholder_format is not None and shape.placeholder_format.type is not None and
                        str(shape.placeholder_format.type).startswith("TITLE")
                    )
                except Exception:
                    pass
                elements.append({"type": "text", "isTitle": is_title, "paragraphs": paragraphs, **meta, **rect})
        except Exception:
            continue
    return elements


def parse_pptx(path, assets_dir, asset_prefix):
    os.makedirs(assets_dir, exist_ok=True)
    prs = Presentation(path)
    slide_w = prs.slide_width
    slide_h = prs.slide_height
    counter_ref = [0]
    theme = _load_theme_map(prs)

    slides_out = []
    for slide in prs.slides:
        bg = _slide_background(slide, prs, theme)
        elements = _walk_shapes(slide.shapes, prs, assets_dir, asset_prefix, counter_ref, theme)
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text
        except Exception:
            notes = ""
        slides_out.append({
            "background": bg,
            "elements": elements,
            "notes": notes,
        })

    return {
        "slideWidthEmu": slide_w,
        "slideHeightEmu": slide_h,
        "aspectRatio": round(slide_w / slide_h, 6) if slide_h else 1.7778,
        "slides": slides_out,
    }
